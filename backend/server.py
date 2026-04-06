"""FastAPI backend — IR data management + PPT export."""

import json
import io
import base64
from pathlib import Path

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response

from schema import IRDocument, SLIDE_NAMES

app = FastAPI(title="Rovothome IR Manager")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

DATA_DIR = Path(__file__).parent / "data"
DATA_FILE = DATA_DIR / "ir_data.json"
OUTPUT_DIR = Path(__file__).parent.parent / "output"
OUTPUT_DIR.mkdir(exist_ok=True)


def _load() -> IRDocument:
    if DATA_FILE.exists():
        raw = json.loads(DATA_FILE.read_text(encoding="utf-8"))
        return IRDocument(**raw)
    return IRDocument()


def _save(doc: IRDocument):
    DATA_DIR.mkdir(exist_ok=True)
    DATA_FILE.write_text(doc.model_dump_json(indent=2), encoding="utf-8")


# --- API endpoints ---

@app.get("/api/slides")
def get_all():
    """Get full IR document data."""
    doc = _load()
    return {
        "data": doc.model_dump(),
        "slide_names": SLIDE_NAMES,
    }


@app.put("/api/slides")
def save_all(doc: IRDocument):
    """Save full IR document."""
    _save(doc)
    return {"ok": True}


@app.patch("/api/slides/{slide_id}")
def update_slide(slide_id: str, body: dict):
    """Update a single slide's data."""
    doc = _load()
    if not hasattr(doc, slide_id):
        return {"error": f"Unknown slide: {slide_id}"}

    current = getattr(doc, slide_id)
    updated = current.model_copy(update=body)
    setattr(doc, slide_id, updated)
    _save(doc)
    return {"ok": True, "slide_id": slide_id}


@app.put("/api/config")
def update_config(config: list[dict]):
    """Update slide order and visibility."""
    from schema import SlideConfig
    doc = _load()
    doc.slide_config = [SlideConfig(**c) for c in config]
    _save(doc)
    return {"ok": True}


@app.post("/api/export")
def export_pptx():
    """Generate and return PPT file."""
    doc = _load()
    from pptx_engine import generate_pptx
    output_path = OUTPUT_DIR / "rovothome_ir.pptx"
    generate_pptx(doc, output_path)
    return FileResponse(
        output_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename="Rovothome_IR.pptx",
    )


@app.get("/api/export/download")
def download_pptx():
    """Download last generated PPT."""
    output_path = OUTPUT_DIR / "rovothome_ir.pptx"
    if not output_path.exists():
        doc = _load()
        from pptx_engine import generate_pptx
        generate_pptx(doc, output_path)
    return FileResponse(output_path, filename="Rovothome_IR.pptx")


@app.get("/api/preview/{slide_id}")
def preview_slide(slide_id: str):
    """Generate a single slide preview as PNG image."""
    doc = _load()
    from pptx_engine import generate_single_slide_png
    img_bytes = generate_single_slide_png(doc, slide_id)
    if img_bytes is None:
        return {"error": "Could not generate preview"}
    b64 = base64.b64encode(img_bytes).decode()
    return {"image": b64, "slide_id": slide_id}


@app.get("/api/financials/summary")
def financial_summary():
    """Get computed financial summary for sync across slides."""
    doc = _load()
    pl = doc.pl
    summary = {}
    for row in pl.rows:
        for year, val in row.values.items():
            key = f"{row.label}_{year}"
            summary[key] = val
    # Also expose top-level numbers
    if pl.rows:
        revenue_row = next((r for r in pl.rows if "매출" in r.label), None)
        if revenue_row:
            summary["latest_revenue"] = list(revenue_row.values.values())[-1] if revenue_row.values else "0"
            summary["revenue_by_year"] = revenue_row.values
    return summary


@app.post("/api/import/financial")
async def import_financial(file: UploadFile = File(...)):
    """Import financial data from Excel or CSV file.

    Expected format:
    - Row headers in first column (매출, 영업이익, etc.)
    - Year headers in first row (2026, 2027, etc.)
    """
    import tempfile
    contents = await file.read()
    suffix = Path(file.filename).suffix.lower()

    try:
        if suffix in (".xlsx", ".xls"):
            import openpyxl
            tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
            tmp.write(contents)
            tmp.close()
            wb = openpyxl.load_workbook(tmp.name, data_only=True)
            ws = wb.active

            # Parse headers (years) from first row
            years = []
            for col in range(2, ws.max_column + 1):
                val = ws.cell(1, col).value
                if val:
                    years.append(str(val).strip())

            # Parse rows
            rows = []
            for row in range(2, ws.max_row + 1):
                label = ws.cell(row, 1).value
                if not label:
                    continue
                values = {}
                for ci, yr in enumerate(years):
                    cell_val = ws.cell(row, ci + 2).value
                    if cell_val is not None:
                        values[yr] = str(cell_val).strip()
                    else:
                        values[yr] = ""
                rows.append({"label": str(label).strip(), "values": values})

        elif suffix == ".csv":
            import csv
            text = contents.decode("utf-8-sig")
            reader = csv.reader(text.splitlines())
            header = next(reader)
            years = [h.strip() for h in header[1:] if h.strip()]
            rows = []
            for row_data in reader:
                if not row_data or not row_data[0].strip():
                    continue
                label = row_data[0].strip()
                values = {}
                for ci, yr in enumerate(years):
                    val = row_data[ci + 1].strip() if ci + 1 < len(row_data) else ""
                    values[yr] = val
                rows.append({"label": label, "values": values})
        else:
            return {"error": f"지원하지 않는 파일 형식: {suffix}. .xlsx 또는 .csv를 사용해주세요."}

        # Update document
        from schema import PLRow
        doc = _load()
        doc.pl.years = years
        doc.pl.rows = [PLRow(**r) for r in rows]
        _save(doc)

        return {
            "ok": True,
            "years": years,
            "rows": rows,
            "message": f"{len(rows)}개 항목, {len(years)}개 연도 데이터를 가져왔습니다.",
        }

    except Exception as e:
        return {"error": f"파일 파싱 실패: {str(e)}"}


@app.post("/api/chart/pl")
def generate_pl_chart():
    """Generate P&L chart as base64 PNG."""
    doc = _load()
    from charts import render_pl_chart
    img_bytes = render_pl_chart(doc.pl)
    b64 = base64.b64encode(img_bytes).decode()
    return {"image": b64}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
