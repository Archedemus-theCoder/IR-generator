"""PPT slide assembly — matching actual IR template style.

Design system:
- White background, light gray (#F2F2F2) card panels
- Orange (#E8600A) accent for section labels and keyword highlights
- Black text with orange highlights for emphasis
- Pattern: orange section label → big bold headline → content below
- Rounded-corner gray panels for grouping
- Slide numbers top-right
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from .brand_manager import hex_to_rgb
from .models import BrandTheme, Slide

W = 13.333  # Slide width in inches (16:9)
H = 7.5     # Slide height in inches
MARGIN = 0.7


def _rgb(hex_color: str) -> RGBColor:
    r, g, b = hex_to_rgb(hex_color)
    return RGBColor(r, g, b)


# --- Low-level helpers ---

def _rect(sld, left, top, w, h, color):
    s = sld.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _rounded_rect(sld, left, top, w, h, color):
    s = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    s.adjustments[0] = 0.03
    return s


def _text_box(sld, left, top, w, h, text, theme,
              size=14, bold=False, color=None, align=PP_ALIGN.LEFT, font=None):
    """Add a text box. Supports **bold** and <<accent>> markup."""
    tb = sld.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True

    lines = text.strip().split("\n") if text.strip() else [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        p.space_before = Pt(1)

        stripped = line.strip()
        if not stripped:
            run = p.add_run()
            run.text = " "
            run.font.size = Pt(6)
            continue

        # Bullet handling
        if stripped.startswith("- ") or stripped.startswith("* "):
            p.level = 0
            stripped = stripped[2:]
        elif stripped.startswith("  - ") or stripped.startswith("  * "):
            p.level = 1
            stripped = stripped[4:]

        # Parse: **bold**, <<accent text>> (orange highlight)
        _add_rich_text(p, stripped, theme, size, bold, color, font)

    return tb


def _add_rich_text(p, text, theme, size, base_bold, base_color, font):
    """Parse **bold** and <<accent>> in text and add runs."""
    # First split on <<...>> for accent
    accent_parts = re.split(r"(<<.*?>>)", text)
    for apart in accent_parts:
        if apart.startswith("<<") and apart.endswith(">>"):
            inner = apart[2:-2]
            # Could still have **bold** inside
            _add_bold_runs(p, inner, theme, size, True, theme.colors.accent, font)
        else:
            _add_bold_runs(p, apart, theme, size, base_bold, base_color, font)


def _add_bold_runs(p, text, theme, size, base_bold, base_color, font):
    """Parse **bold** markup and add runs."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        run = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.font.bold = True
        else:
            run.text = part
            if base_bold:
                run.font.bold = True
        run.font.size = Pt(size)
        run.font.name = font or theme.fonts.body
        run.font.color.rgb = _rgb(base_color or theme.colors.text_primary)


def _slide_number(sld, num, theme):
    """Slide number in top-right corner."""
    _text_box(sld, W - 1.0, 0.3, 0.5, 0.4, str(num), theme,
              size=theme.sizes.slide_number, color=theme.colors.text_secondary,
              align=PP_ALIGN.RIGHT)


def _section_label(sld, text, theme, left=MARGIN, top=0.4):
    """Orange section label (small, bold, top-left)."""
    _text_box(sld, left, top, 8, 0.4, text, theme,
              size=theme.sizes.section_label, bold=True,
              color=theme.colors.accent, font=theme.fonts.heading)


def _headline(sld, text, theme, left=MARGIN, top=0.85, width=None):
    """Large bold headline with <<accent>> support."""
    w = width or (W - MARGIN * 2)
    _text_box(sld, left, top, w, 1.5, text, theme,
              size=theme.sizes.title, bold=True,
              color=theme.colors.primary, font=theme.fonts.heading)


def _card_header(sld, text, theme, left, top, width):
    """Black header bar inside a card."""
    _rect(sld, left, top, width, 0.45, theme.colors.header_bg)
    _text_box(sld, left + 0.15, top + 0.05, width - 0.3, 0.35, text, theme,
              size=12, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)


def _body_text(sld, text, theme, left, top, width, height):
    """Standard body text with markdown parsing."""
    body = _strip_md_headers(text)
    _text_box(sld, left, top, width, height, body, theme,
              size=theme.sizes.body, color=theme.colors.primary)


# --- Common slide header pattern ---

def _slide_header(sld, slide, theme, slide_num):
    """Standard header: section label + headline + slide number.

    Convention:
    - front_matter.title = section label (orange, small)
    - First ## line in body = headline (large, bold)
    - Rest = body content
    """
    _slide_number(sld, slide_num, theme)

    # Section label
    if slide.front_matter.title:
        _section_label(sld, slide.front_matter.title, theme)

    # Extract headline (first ## line) and remaining body
    headline, body = _extract_headline(slide.rendered_body)

    if headline:
        _headline(sld, headline, theme)

    return headline, body


def _extract_headline(rendered_body):
    """Split body into headline (first ## line) and remaining body."""
    lines = rendered_body.strip().split("\n")
    headline = ""
    body_start = 0

    for i, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("## ") or stripped.startswith("# "):
            headline = re.sub(r"^#+\s*", "", stripped)
            body_start = i + 1
            # Check for continuation lines (non-blank, non-bullet, non-header lines)
            for j in range(i + 1, len(lines)):
                next_line = lines[j].strip()
                if next_line and not next_line.startswith(("#", "-", "*")) and not next_line.startswith("  "):
                    headline += "\n" + next_line
                    body_start = j + 1
                else:
                    break
            break

    remaining = "\n".join(lines[body_start:]).strip()
    return headline, remaining


def _strip_md_headers(text):
    lines = []
    for line in text.split("\n"):
        lines.append(re.sub(r"^#+\s*", "", line))
    return "\n".join(lines)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# =====================================================
# SLIDE BUILDERS
# =====================================================

def build_title_slide(prs, slide, theme, context, **kw):
    """Cover slide: full dark background, centered title."""
    sld = _blank(prs)
    bg = sld.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(theme.colors.primary)

    # Top accent bar
    _rect(sld, 0, 0, W, 0.12, theme.colors.accent)

    # Title centered
    title = slide.front_matter.title
    _text_box(sld, 1.5, 2.0, W - 3, 2.0, title, theme,
              size=40, bold=True, color="#FFFFFF",
              align=PP_ALIGN.CENTER, font=theme.fonts.heading)

    # Subtitle
    if slide.rendered_body:
        first = re.sub(r"^#+\s*", "", slide.rendered_body.strip().split("\n")[0])
        _text_box(sld, 2, 4.2, W - 4, 1.0, first, theme,
                  size=18, color="#FFFFFF", align=PP_ALIGN.CENTER)

    # Bottom accent bar
    _rect(sld, 0, H - 0.08, W, 0.08, theme.colors.accent)


def build_content_slide(prs, slide, theme, context, chart_path=None, **kw):
    """Standard content: section label → headline → body text, optional chart."""
    sld = _blank(prs)
    num = len(prs.slides)
    headline, body = _slide_header(sld, slide, theme, num)

    body_top = 2.5 if headline else 1.5

    if chart_path and chart_path.exists():
        # Text left, chart right
        _body_text(sld, body, theme, MARGIN, body_top, 5.5, 4.5)
        _rounded_rect(sld, 6.8, body_top, 5.8, 4.5, theme.colors.card_bg)
        sld.shapes.add_picture(
            str(chart_path), Inches(7.0), Inches(body_top + 0.15),
            Inches(5.4), Inches(4.2),
        )
    else:
        _body_text(sld, body, theme, MARGIN, body_top, W - MARGIN * 2, 4.5)


def build_two_column_slide(prs, slide, theme, context, chart_path=None, **kw):
    """Text left + chart/metric cards right."""
    sld = _blank(prs)
    num = len(prs.slides)
    headline, body = _slide_header(sld, slide, theme, num)

    body_top = 2.5 if headline else 1.5

    # Left: body text
    _body_text(sld, body, theme, MARGIN, body_top, 5.5, 4.5)

    # Right: chart or auto metric cards
    right_left = 6.8
    right_w = 5.8

    if chart_path and chart_path.exists():
        _rounded_rect(sld, right_left, body_top, right_w, 4.5, theme.colors.card_bg)
        sld.shapes.add_picture(
            str(chart_path), Inches(right_left + 0.15), Inches(body_top + 0.15),
            Inches(right_w - 0.3), Inches(4.2),
        )
    else:
        _build_auto_metrics(sld, theme, context, right_left, body_top, right_w)


def build_two_panel_slide(prs, slide, theme, context, **kw):
    """Two equal side-by-side panels (comparison layout)."""
    sld = _blank(prs)
    num = len(prs.slides)
    headline, body = _slide_header(sld, slide, theme, num)

    body_top = 2.5 if headline else 1.5
    panel_w = 5.8
    panel_h = 4.3
    gap = 0.4

    # Parse body into two sections (split by ---)
    sections = body.split("---")
    left_text = sections[0].strip() if len(sections) >= 1 else ""
    right_text = sections[1].strip() if len(sections) >= 2 else ""

    # Left panel
    left_x = MARGIN
    _rounded_rect(sld, left_x, body_top, panel_w, panel_h, theme.colors.card_bg)

    left_lines = left_text.split("\n", 1)
    left_title = re.sub(r"^#+\s*", "", left_lines[0]) if left_lines else ""
    left_body = left_lines[1].strip() if len(left_lines) > 1 else ""

    if left_title:
        _card_header(sld, left_title, theme, left_x, body_top, panel_w)
        _body_text(sld, left_body, theme, left_x + 0.3, body_top + 0.6, panel_w - 0.6, panel_h - 0.9)
    else:
        _body_text(sld, left_text, theme, left_x + 0.3, body_top + 0.3, panel_w - 0.6, panel_h - 0.6)

    # Right panel
    right_x = MARGIN + panel_w + gap
    _rounded_rect(sld, right_x, body_top, panel_w, panel_h, theme.colors.card_bg)

    right_lines = right_text.split("\n", 1)
    right_title = re.sub(r"^#+\s*", "", right_lines[0]) if right_lines else ""
    right_body = right_lines[1].strip() if len(right_lines) > 1 else ""

    if right_title:
        _card_header(sld, right_title, theme, right_x, body_top, panel_w)
        _body_text(sld, right_body, theme, right_x + 0.3, body_top + 0.6, panel_w - 0.6, panel_h - 0.9)
    else:
        _body_text(sld, right_text, theme, right_x + 0.3, body_top + 0.3, panel_w - 0.6, panel_h - 0.6)

    _slide_number(sld, num, theme)


def build_three_column_slide(prs, slide, theme, context, **kw):
    """Three cards side by side."""
    sld = _blank(prs)
    num = len(prs.slides)
    headline, body = _slide_header(sld, slide, theme, num)

    body_top = 2.5 if headline else 1.5
    card_w = 3.7
    card_h = 4.3
    gap = 0.3

    sections = body.split("---")
    total_w = 3 * card_w + 2 * gap
    start_x = (W - total_w) / 2

    accent_colors = [theme.colors.accent, theme.colors.primary, theme.colors.positive]

    for idx in range(3):
        text = sections[idx].strip() if idx < len(sections) else ""
        card_x = start_x + idx * (card_w + gap)

        _rounded_rect(sld, card_x, body_top, card_w, card_h, theme.colors.card_bg)

        # Top accent line
        _rect(sld, card_x + 0.15, body_top + 0.12, card_w - 0.3, 0.06, accent_colors[idx % 3])

        lines = text.split("\n", 1)
        card_title = re.sub(r"^#+\s*", "", lines[0]) if lines else ""
        card_body = lines[1].strip() if len(lines) > 1 else ""

        # Card title
        _text_box(sld, card_x + 0.25, body_top + 0.35, card_w - 0.5, 0.5,
                  card_title, theme, size=16, bold=True,
                  color=theme.colors.primary, font=theme.fonts.heading)

        # Card body
        if card_body:
            _body_text(sld, card_body, theme, card_x + 0.25, body_top + 0.95,
                       card_w - 0.5, card_h - 1.2)

    _slide_number(sld, num, theme)


def build_metric_highlight_slide(prs, slide, theme, context, **kw):
    """Big metric number with KPI cards below."""
    sld = _blank(prs)
    num = len(prs.slides)
    headline, body = _slide_header(sld, slide, theme, num)

    body_lines = body.strip().split("\n") if body else []
    metric_text = body_lines[0].strip() if body_lines else ""
    supporting = [l.strip().lstrip("- *") for l in body_lines[1:] if l.strip()]

    # Big metric
    metric_top = 2.2 if headline else 1.5
    _text_box(sld, 1, metric_top, W - 2, 1.5, metric_text, theme,
              size=theme.sizes.metric_highlight, bold=True,
              color=theme.colors.accent, align=PP_ALIGN.CENTER,
              font=theme.fonts.heading)

    # KPI cards
    if supporting:
        n = min(len(supporting), 4)
        card_w = 2.6
        gap = 0.25
        total_w = n * card_w + (n - 1) * gap
        start_x = (W - total_w) / 2
        card_top = metric_top + 2.0

        accent_colors = [theme.colors.accent, theme.colors.primary, theme.colors.positive, "#6366F1"]

        for idx, item in enumerate(supporting[:4]):
            cx = start_x + idx * (card_w + gap)
            _rounded_rect(sld, cx, card_top, card_w, 2.0, theme.colors.card_bg)
            _rect(sld, cx + 0.1, card_top + 0.08, card_w - 0.2, 0.05, accent_colors[idx % 4])

            # Try to split "value label"
            parts = item.split(" ", 1)
            if parts[0] and any(c.isdigit() or c == '%' for c in parts[0]):
                _text_box(sld, cx + 0.15, card_top + 0.3, card_w - 0.3, 0.7,
                          parts[0], theme, size=26, bold=True,
                          color=theme.colors.primary, align=PP_ALIGN.CENTER,
                          font=theme.fonts.heading)
                if len(parts) > 1:
                    _text_box(sld, cx + 0.15, card_top + 1.1, card_w - 0.3, 0.7,
                              parts[1], theme, size=11,
                              color=theme.colors.text_secondary, align=PP_ALIGN.CENTER)
            else:
                _text_box(sld, cx + 0.15, card_top + 0.5, card_w - 0.3, 1.2,
                          item, theme, size=13,
                          color=theme.colors.primary, align=PP_ALIGN.CENTER)

    _slide_number(sld, num, theme)


def build_table_slide(prs, slide, theme, context, **kw):
    """Financial table with styled header."""
    sld = _blank(prs)
    num = len(prs.slides)
    _slide_header(sld, slide, theme, num)

    periods = context.get("periods", [])[:8]
    revenue = context.get("revenue_by_period", [])[:8]
    costs = context.get("costs_by_period", [])[:8]
    profit = context.get("profit_by_period", [])[:8]

    if not periods:
        build_content_slide(prs, slide, theme, context)
        return

    rows = 4
    cols = len(periods) + 1

    tbl = sld.shapes.add_table(
        rows, cols, Inches(MARGIN), Inches(2.5),
        Inches(W - MARGIN * 2), Inches(3.5),
    )
    table = tbl.table

    # Header
    for i in range(cols):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(theme.colors.header_bg)
        text = "" if i == 0 else periods[i - 1]
        _style_cell(cell, text, theme, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)

    # Data
    data_rows = [
        ("Revenue", revenue, theme.colors.accent),
        ("Costs", costs, theme.colors.primary),
        ("Profit", profit, None),
    ]

    for row_idx, (label, values, color) in enumerate(data_rows, 1):
        _style_cell(table.cell(row_idx, 0), label, theme, bold=True)

        # Alternating row bg
        if row_idx % 2 == 0:
            for c in range(cols):
                table.cell(row_idx, c).fill.solid()
                table.cell(row_idx, c).fill.fore_color.rgb = _rgb(theme.colors.card_bg)

        for col_idx, val in enumerate(values):
            text = _fmt_num(val)
            c = color
            if label == "Profit":
                c = theme.colors.positive if val >= 0 else theme.colors.negative
            _style_cell(table.cell(row_idx, col_idx + 1), text, theme, color=c)


def build_section_divider_slide(prs, slide, theme, context, **kw):
    """Section divider: gray background with centered section title."""
    sld = _blank(prs)
    num = len(prs.slides)

    bg = sld.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(theme.colors.card_bg)

    # Accent bar
    _rect(sld, 0, 0, W, 0.08, theme.colors.accent)

    # Section title centered
    title = slide.front_matter.title
    _text_box(sld, 2, 2.5, W - 4, 2.0, title, theme,
              size=36, bold=True, color=theme.colors.primary,
              align=PP_ALIGN.CENTER, font=theme.fonts.heading)

    # Subtitle from body
    if slide.rendered_body:
        first = re.sub(r"^#+\s*", "", slide.rendered_body.strip().split("\n")[0])
        _text_box(sld, 2, 4.5, W - 4, 1.0, first, theme,
                  size=16, color=theme.colors.text_secondary,
                  align=PP_ALIGN.CENTER)

    _slide_number(sld, num, theme)


# --- Auto metric cards (when no chart) ---

def _build_auto_metrics(sld, theme, context, left, top, width):
    """Auto-generate metric cards from financial context."""
    years = sorted(set(
        k.split("_")[-1]
        for k in context
        if k.startswith("total_revenue_") and k.split("_")[-1].isdigit()
    ))
    if not years:
        return

    latest = years[-1]
    rev = context.get(f"total_revenue_{latest}", 0)
    margin = context.get(f"gross_margin_{latest}", 0)
    profit = context.get(f"gross_profit_{latest}", 0)

    metrics = [
        (_fmt_compact(rev), f"FY{latest} Revenue", theme.colors.accent),
        (f"{margin*100:.0f}%", f"FY{latest} Gross Margin", theme.colors.primary),
        (_fmt_compact(profit), f"FY{latest} Profit", theme.colors.positive if profit >= 0 else theme.colors.negative),
    ]

    card_h = 1.3
    gap = 0.2

    for idx, (value, label, color) in enumerate(metrics):
        cy = top + idx * (card_h + gap)
        _rounded_rect(sld, left, cy, width, card_h, theme.colors.card_bg)
        _rect(sld, left, cy + 0.15, 0.07, card_h - 0.3, color)
        _text_box(sld, left + 0.3, cy + 0.1, width - 0.5, 0.6,
                  value, theme, size=24, bold=True, color=theme.colors.primary,
                  font=theme.fonts.heading)
        _text_box(sld, left + 0.3, cy + 0.75, width - 0.5, 0.4,
                  label, theme, size=11, color=theme.colors.text_secondary)


# --- Cell helpers ---

def _style_cell(cell, text, theme, bold=False, color=None, align=PP_ALIGN.RIGHT):
    cell.text = text
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = theme.fonts.body
            run.font.bold = bold
            if color:
                run.font.color.rgb = _rgb(color)


def _fmt_num(v):
    if abs(v) >= 1_000_000:
        return f"${v/1_000_000:.1f}M"
    if abs(v) >= 1_000:
        return f"${v/1_000:,.0f}K"
    return f"${v:,.0f}"


def _fmt_compact(v):
    if abs(v) >= 1_000_000_000:
        return f"${v/1_000_000_000:.1f}B"
    if abs(v) >= 1_000_000:
        return f"${v/1_000_000:.1f}M"
    if abs(v) >= 1_000:
        return f"${v/1_000:.0f}K"
    return f"${v:,.0f}"


# --- Registry ---

LAYOUT_BUILDERS = {
    "title": build_title_slide,
    "content": build_content_slide,
    "two_column": build_two_column_slide,
    "two_panel": build_two_panel_slide,
    "three_column": build_three_column_slide,
    "metric_highlight": build_metric_highlight_slide,
    "table": build_table_slide,
    "section_divider": build_section_divider_slide,
}


def create_presentation(slides, theme, context, chart_paths,
                        base_template=None, output_path=Path("output/deck.pptx")):
    if base_template and base_template.exists():
        prs = Presentation(str(base_template))
    else:
        prs = Presentation()
        prs.slide_width = Inches(W)
        prs.slide_height = Inches(H)

    for slide in slides:
        layout = slide.front_matter.layout
        builder = LAYOUT_BUILDERS.get(layout, build_content_slide)
        chart_path = chart_paths.get(slide.front_matter.chart) if slide.front_matter.chart else None
        builder(prs, slide, theme, context, chart_path=chart_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path
